#!/usr/bin/env python3
"""Build GOST-formatted lab report from Markdown.

Usage:
    python build.py <report.md> [output.docx]

Expects templates/ and filters/ directories next to scripts/.
"""

import shutil
import subprocess
import sys
import tempfile
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docxcompose.composer import Composer

SKILL_DIR = Path(__file__).resolve().parent.parent
TEMPLATES_DIR = SKILL_DIR / "templates"
FILTERS_DIR = SKILL_DIR / "filters"

REFERENCE_DOC = TEMPLATES_DIR / "reference.docx"
PAGEBREAK_FILTER = FILTERS_DIR / "pagebreak.lua"

CAPS_FIELDS = {"LAB_TITLE", "DISCIPLINE"}


def parse_yaml_front_matter(md_path: Path) -> dict:
    text = md_path.read_text(encoding="utf-8")
    match = re.match(r"^---\s*\n(.*?)\n---", text, re.DOTALL)
    if not match:
        return {}
    metadata = {}
    for line in match.group(1).splitlines():
        line = line.strip()
        if ":" in line:
            key, _, value = line.partition(":")
            key = key.strip()
            value = value.strip().strip('"').strip("'")
            metadata[key] = value
    return metadata


def fill_title_page(src_path: Path, metadata: dict, output_path: Path):
    shutil.copy2(src_path, output_path)
    doc = Document(str(output_path))

    placeholders = {
        "{{TEACHER_TITLE}}": metadata.get("teacher_title", ""),
        "{{TEACHER_NAME}}": metadata.get("teacher_name", ""),
        "{{LAB_NUMBER}}": metadata.get("lab_number", ""),
        "{{LAB_TITLE}}": metadata.get("lab_title", ""),
        "{{DISCIPLINE}}": metadata.get("discipline", ""),
        "{{GROUP}}": metadata.get("group", ""),
        "{{STUDENT_NAME}}": metadata.get("student_name", ""),
        "{{DEPARTMENT}}": metadata.get("department", ""),
        "{{YEAR}}": str(datetime.now().year),
    }

    for key in list(placeholders):
        field_name = key.strip("{}")
        if field_name in CAPS_FIELDS:
            placeholders[key] = placeholders[key].upper()

    def replace_in_runs(runs):
        for run in runs:
            for placeholder, value in placeholders.items():
                if placeholder in run.text:
                    run.text = run.text.replace(placeholder, value)

    for paragraph in doc.paragraphs:
        replace_in_runs(paragraph.runs)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_runs(paragraph.runs)

    doc.save(str(output_path))


def preprocess_markdown(md_path: Path, output_path: Path):
    text = md_path.read_text(encoding="utf-8")
    text = text.replace("\u2014", "\u2013")  # em dash -> en dash
    text = re.sub(r"(?m)^title:.*\n", "", text)  # strip title from YAML (already on title page)
    output_path.write_text(text, encoding="utf-8")


def build_report_body(md_path: Path, output_path: Path):
    preprocessed = md_path.parent / f".{md_path.stem}_preprocessed.md"
    preprocess_markdown(md_path, preprocessed)
    try:
        cmd = [
            "pandoc",
            str(preprocessed),
            "-o",
            str(output_path),
            f"--reference-doc={REFERENCE_DOC}",
            f"--lua-filter={PAGEBREAK_FILTER}",
        ]
        subprocess.run(cmd, check=True, cwd=md_path.parent)
    finally:
        preprocessed.unlink(missing_ok=True)


def merge_documents(title_path: Path, body_path: Path, output_path: Path):
    title_doc = Document(str(title_path))
    composer = Composer(title_doc)
    body_doc = Document(str(body_path))
    composer.append(body_doc)
    composer.save(str(output_path))


def main():
    if len(sys.argv) < 2:
        print("Usage: python build.py <report.md> [output.docx]")
        sys.exit(1)

    if not shutil.which("pandoc"):
        print("Error: pandoc not found. Install it: https://pandoc.org/installing.html")
        sys.exit(1)

    md_path = Path(sys.argv[1]).resolve()
    if not md_path.exists():
        print(f"File not found: {md_path}")
        sys.exit(1)

    output_path = (
        Path(sys.argv[2]).resolve() if len(sys.argv) >= 3 else md_path.with_suffix(".docx")
    )

    metadata = parse_yaml_front_matter(md_path)
    print(f"Metadata: {metadata}")

    with tempfile.TemporaryDirectory() as tmpdir:
        tmpdir = Path(tmpdir)

        filled = tmpdir / "filled.docx"
        fill_title_page(REFERENCE_DOC, metadata, filled)
        print("Title page filled")

        body_docx = tmpdir / "body.docx"
        build_report_body(md_path, body_docx)
        print("Report body built")

        merge_documents(filled, body_docx, output_path)
        print(f"Report saved: {output_path}")


if __name__ == "__main__":
    main()
